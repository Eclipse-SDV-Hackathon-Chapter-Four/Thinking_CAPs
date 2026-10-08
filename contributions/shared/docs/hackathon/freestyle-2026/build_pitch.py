#!/usr/bin/env python3
"""Build the editable Freestyle pitch from the official Eclipse template.

Dependencies: python-pptx, Pillow. Run from any directory.
PDF export: libreoffice --headless --convert-to pdf --outdir <this folder> <pptx>.
"""
from copy import deepcopy
from pathlib import Path
import json

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / 'assets'
REPO = 'https://github.com/Eclipse-SDV-Hackathon-Chapter-Four/Thinking_CAPs'
BASE = REPO + '/blob/38fbb58b4420365a8c02de1552ed983d51e9a499/'
TEMPLATE = 'https://docs.google.com/presentation/d/1LxFKeG0z-uxOMgOhZX4OWIPQG5Vmeo1Tptall_lJ0_Y/edit'
GUIDE = 'https://github.com/Eclipse-SDV-Hackathon-Chapter-Four/.github/blob/main/profile/SDV%20Hackathon%202026_Pitching%20Session.pdf'
RUBRIC = 'https://github.com/Eclipse-SDV-Hackathon-Chapter-Four/.github/blob/main/profile/Eclipse_SDV_Hackathon_2026_EvaluationForms.pdf'
PURPLE, TEAL, INK, MUTED, LIGHT = '701C7F', '0097A7', '222222', '595959', 'F4EEF5'
prs = Presentation(ASSETS / 'official-template.pptx')
plan = []
notes = []

def style(run, size=18, color=INK, bold=False, font='Roboto'):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = False
    run.font.color.rgb = RGBColor.from_string(color)

def text(shape, value, size=18, color=INK, bold=False, font='Roboto'):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.02)
    tf.margin_top = tf.margin_bottom = Inches(.02)
    for i, line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(5)
        pPr = p._p.get_or_add_pPr()
        # clear() preserves paragraph properties. Replace the complete bullet
        # group, including an inherited buNone, instead of appending duplicates.
        # Keep the no-bullet choice in DrawingML schema order.
        for el in list(pPr):
            if el.tag.rsplit('}', 1)[-1].startswith('bu'):
                pPr.remove(el)
        pPr.insert_element_before(OxmlElement('a:buNone'), 'a:tabLst', 'a:defRPr', 'a:extLst')
        style(p.add_run(), size, color, bold, font)
        p.runs[-1].text = line
    return shape

def box(slide, x, y, w, h, value='', size=18, color=INK, bold=False, font='Roboto'):
    return text(slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)), value, size, color, bold, font)

def rect(slide, x, y, w, h, color=LIGHT):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(color)
    sh.line.fill.background()
    return sh

def link(slide, x, y, w, label, url):
    sh = box(slide, x, y, w, .23, label, 9, TEAL)
    sh.text_frame.paragraphs[0].runs[0].hyperlink.address = url
    return sh

def fit_image(slide, name, x, y, w, h):
    path = ASSETS / name
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    ow, oh = iw * scale, ih * scale
    slide.shapes.add_picture(str(path), Inches(x+(w-ow)/2), Inches(y+(h-oh)/2), Inches(ow), Inches(oh))

def clone(src):
    dst = prs.slides.add_slide(src.slide_layout)
    for sh in list(dst.shapes):
        sh._element.getparent().remove(sh._element)
    relmap = {}
    for rel in src.part.rels.values():
        if rel.reltype.endswith(('/slideLayout', '/notesSlide')):
            continue
        relmap[rel.rId] = dst.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref) if rel.is_external else dst.part.rels.get_or_add(rel.reltype, rel.target_part)
    for sh in src.shapes:
        el = deepcopy(sh._element)
        for node in el.iter():
            for attr, val in list(node.attrib.items()):
                if attr.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}') and val in relmap:
                    node.set(attr, relmap[val])
        dst.shapes._spTree.insert_element_before(el, 'p:extLst')
    if src._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}cSld/{http://schemas.openxmlformats.org/presentationml/2006/main}bg') is not None:
        dst._element.cSld.insert(0, deepcopy(src._element.cSld.bg))
    return dst

def register(slide, title, criterion, weight, seconds, speech, sources, exemplar):
    slide.notes_slide.notes_text_frame.text = speech + '\n\nEvidence:\n' + '\n'.join(label + ': ' + url for label, url in sources)
    item = dict(title=title, criterion=criterion, weight_percent=weight, seconds=seconds,
                template_exemplar=exemplar, media_mapping='Keep Eclipse branding and footer; replace instructional text. Added evidence media is explicitly labelled.',
                sources=[dict(label=l, url=u) for l,u in sources])
    plan.append(item); notes.append((title, criterion, seconds, speech, sources))

def content(slide, title, lead, criterion, weight, seconds, speech, sources, exemplar):
    # Preserve and populate the exemplar's native title/body objects and footer.
    native = [s for s in slide.shapes if s.has_text_frame]
    title_sh, body_sh = native[-1], native[-2]
    text(title_sh, title, 18.2, PURPLE, True, 'Inter')
    title_sh.width = Inches(8.55)
    text(body_sh, lead, 18)
    body_sh.height = Inches(.7)
    box(slide, .8, .34, 8.5, .24, criterion.upper() + (f'  |  {weight}%' if weight else ''), 9, TEAL, True, 'Inter')
    box(slide, 8.1, 5.27, 1.25, .2, 'THINKING CAPs', 7, MUTED, True, 'Inter')
    for i,(label,url) in enumerate(sources[:3]):
        link(slide, .8+i*2.85, 4.85, 2.75, label, url)
    register(slide,title,criterion,weight,seconds,speech,sources,exemplar)
    return body_sh

def columns(slide, items, y=2.12, height=2.28):
    for i,(label,headline,body) in enumerate(items):
        x=.8+i*2.85
        rect(slide,x,y,2.7,height)
        box(slide,x+.14,y+.12,2.4,.27,label,10,TEAL,True,'Inter')
        box(slide,x+.14,y+.5,2.4,.72,headline,18,PURPLE,True,'Inter')
        box(slide,x+.14,y+1.29,2.4,height-1.34,body,14)

def banner(slide, value, y=4.46):
    box(slide,.8,y,8.5,.32,value,11,MUTED)

# Retain the cover's native media, event decoration and mixed hierarchy.
s=prs.slides[0]
text(next(x for x in s.shapes if x.shape_id==61),'Thinking CAPs',28,'FFFFFF',False,'Arial')
text(next(x for x in s.shapes if x.shape_id==70),'Drive. Diagnose.\nVerify. Contribute.',18,'FFFFFF',True,'Inter')
box(s,.8,3.98,5.0,.25,'Freestyle / HackFest  |  Eclipse SDV Hackathon 2026',10,'FFFFFF')
register(s,'Thinking CAPs','Opening',None,30,
    'We are Thinking CAPs: five contributors working across vehicle architecture, middleware, diagnostics, simulation and testing. Our goal is to make Eclipse SDV integration observable and leave useful contributions that the community can review and continue. We will show our upstream changes, their validation, and the connected vehicle evidence. Some core work was prepared on 4 October; the hardware and gateway work includes 6–7 October results. The repository separates prepared assets from event work.',
    [('Team repository',REPO),('Prepared-work inventory',BASE+'contributions/shared/docs/prepared-work.md')],'h6b2788cdbe65c0ea_1_0')

s=prs.slides[1]
content(s,'Existing needs. Reviewable contributions.','We advance project issues with working code, tests and reusable integration.',
 'Contribution Value',25,90,
 'Our contribution portfolio starts with existing project needs. Communication issue 1167 lacked dedicated integration coverage for repeated COM API calls. PR 1335 adds that coverage without changing production APIs. Diagnostics issue 16 needs S-CORE diagnostic resources exposed through OpenSOVD; PR 40 adds the provider adapter. OpenBSW lacked logical-address routing between DoIP and CAN targets; our local transportRouter contribution adds it with tests and a maintainer-facing packet. These are three distinct reviewable scopes. Submitted means open for upstream review, not accepted or merged. Lifecycle configuration deduplication, the SOME/IP registration fix and the fault-storage write-through patch are supporting contributions, with separate records.',
 [('COM PR #1335','https://github.com/eclipse-score/communication/pull/1335'),('Diagnostics PR #40','https://github.com/eclipse-score/inc_diagnostics/pull/40'),('OpenBSW contribution',BASE+'contributions/eclipse-openbsw/transport-router/README.md')], 'h6b2788cdbe65c0ea_1_82')
columns(s,[('SUBMITTED','COM reliability','Repeated offer, stop, discovery and subscription integration tests.'),('SUBMITTED','S-CORE to SOVD','Expose diagnostic data resources through an OpenSOVD provider.'),('LOCAL / VERIFIED','Diagnostic routing','Route UDS by logical address from DoIP to CAN targets.')])
banner(s,'Supporting work: lifecycle #704, SOME/IP #84 and native fault-storage persistence.')

s=prs.slides[2]
content(s,'Measured quality, with explicit boundaries.','Native tests, source identities and review packets accompany the selected changes.',
 'Technical Quality & Maturity',25,90,
 'The numbers here are retained measurements on explicitly recorded baselines, not new executions during presentation creation. Communication 1167 records 503 passing tests and six skipped platform or configuration targets. Its measured copyright result retains 204 inherited findings with no added findings; these still need upstream disposition. Lifecycle 704 records 113 passing native cases. The OpenBSW contribution packet records 46 passing unit tests, two Bazel tests and 28 passing host gateway integration tests. Its larger application report still has a bus-load failure and unrun integration campaigns. Our software-factory workflows produce patches and evidence; deterministic checks decide results and human acceptance remains explicit. PR 40 still needs author/ECA resolution and native published-revision validation. We do not infer production qualification, full platform coverage or upstream acceptance from these results.',
 [('COM validation',BASE+'contributions/eclipse-score/communication/1167/README.md'),('Lifecycle #704',BASE+'contributions/eclipse-score/lifecycle/704/README.md'),('OpenBSW native gates',BASE+'contributions/eclipse-openbsw/transport-router/validation.md')], 'h6b2788cdbe65c0ea_1_136')
columns(s,[('COMMUNICATION #1167','503 passed','6 skipped; inherited copyright findings retained.'),('LIFECYCLE #704','113 passed','Native cases for shared configuration generation.'),('OPENBSW ROUTER','46 unit tests','2 Bazel tests and 28 host integration tests passed.')])
banner(s,'Remaining gates: author/ECA where applicable, native CI, maintainer review and platform coverage.')

s=prs.slides[3]
content(s,'Connect projects through observable behavior.','Reusable integration paths link vehicle state, receiver diagnosis and embedded actuation.',
 'Eclipse SDV Ecosystem Impact',20,90,
 'Follow the first integration: CARLA and the existing X-Verse vehicle bridge feed the actual S-CORE receiver. The openDuT-managed path carries the selected SOME/IP traffic. Receiver instrumentation and a separate OpenSOVD provider expose the accepted state, freshness and native fault history through a labelled App-data fallback. The project campaign controls disturbances; we do not claim native VIPER execution. In the separate lighting integration, X-Verse carries vehicle status to actual ThreadX on an AZ3166 board using UART/SLCAN, and the board decision returns to CARLA lamps. AutoSD is a separately verified guest target for the Linux ThreadX lighting runtime, bridge and native OpenSOVD provider. These are distinct tested configurations, not a claim that all components form one simultaneously validated chain. The reusable value is the contracts, deployment patterns and evidence linking observations to behavior.',
 [('Receiver campaign',BASE+'contributions/shared/docs/reproduction.md'),('Physical ThreadX',BASE+'demo/X-Verse/external_hackathon_ecus/ThreadX/az3166/README.md'),('AutoSD guest evidence',BASE+'contributions/eclipse-autosd/AutoSD/artifacts/README.md')], 'h6b2788cdbe65c0ea_1_143')
for yy,label,nodes in [(2.12,'RECEIVER / DIAGNOSIS',['CARLA + X-Verse','S-CORE receiver','OpenSOVD']), (3.40,'ZONAL LIGHTING',['X-Verse status','ThreadX AZ3166','CARLA lamps'])]:
    box(s,.8,yy,8.4,.25,label,10,TEAL,True,'Inter')
    for i,name in enumerate(nodes):
        x=.8+i*2.85
        rect(s,x,yy+.36,2.55,.6)
        sh=box(s,x+.08,yy+.47,2.4,.36,name,16,PURPLE,True,'Inter')
        if i<2: box(s,x+2.60,yy+.46,.25,.35,'→',18,TEAL)
box(s,.9,3.05,8.0,.28,'openDuT-managed SOME/IP path; receiver freshness and fault-history observations.',11,MUTED)
banner(s,'AutoSD: separately verified lighting deployment, native diagnosis and disturbance/recovery.')

s=prs.slides[4]
content(s,'A contributor can pick up the next step.','Pinned inputs and portable records make the work understandable and repeatable.',
 'Reusability & Maintainability',10,60,
 'Our handover includes source and tool pins, configuration, manifests, expected verdicts, teardown instructions, and known limitations. A contributor can inspect the patch, rebuild on the documented baseline, run the selected campaign and compare results. The current reference-host instructions still require adapting local paths and acquiring some cached images. Source-clean self-reproduction shares this host, images and LLVM; independent human and cold-machine reproduction remain unattested. The contribution hash verifier currently finds a working-tree mismatch in the 1261 README, which must be reconciled on the final submission revision. AutoSD artifact and current-source verification passed during this review. Hash integrity proves identity, not functional correctness. We provide both patch-level follow-up and integration-level continuation.',
 [('Reproduction guide',BASE+'contributions/shared/docs/reproduction.md'),('Contribution registry',BASE+'contributions/README.md'),('AutoSD verification',BASE+'contributions/eclipse-autosd/AutoSD/artifacts/README.md')], 'h6b2788cdbe65c0ea_1_157')
columns(s,[('01 / IDENTIFY','Pinned inputs','Source revisions, dependencies, configuration and executable hashes.'),('02 / EXECUTE','Repeatable checks','Native regressions, campaign verdicts and owned teardown.'),('03 / CONTINUE','Explicit handover','Patches, PR links, known gaps and maintainer-facing next steps.')])
banner(s,'Reference-host self-run demonstrated; independent operator and cold-machine reproduction pending.')

s=prs.slides[5]
body=content(s,'One failure story. Traceable evidence.','Show what the receiver consumed, what failed, and what recovered.',
 'Pitch & Handover Clarity',8,60,
 'For the demonstration, show four transitions: fresh accepted receiver data; managed communication interruption; stale input and activated fault history while diagnosis remains reachable; then restored traffic, recovery and retained history. Use a current live run only after the environment is ready. If it is not ready, identify the saved physical replay as a recorded 4 October run. The image on this slide is historical evidence from that recorded native campaign. Fault history uses a labelled App-data fallback, not native faults routing. Monitor recovery does not imply automatic cruise re-engagement. The screenshot does not claim exposed control output or E2E sample integrity. After the flow, open a selected PR or patch and its native test record so the handover is concrete.',
 [('Physical verdicts',BASE+'contributions/shared/evidence/f009-reproduction-physical/results.json'),('Dashboard evidence',BASE+'contributions/shared/evidence/f010-live/verification.json'),('Saved physical replay',REPO+'/tree/38fbb58b4420365a8c02de1552ed983d51e9a499/evidence/f009-recorded-replay-final')], 'h6b2788cdbe65c0ea_1_150')
box(s,.8,2.12,4.35,2.25,'01  Fresh receiver observations\n02  Managed communication loss\n03  Fault activation and diagnosis\n04  Recovery with history retained',18)
fit_image(s,'receiver-diagnosis.png',5.5,1.92,3.5,2.58)
box(s,5.35,4.53,3.85,.22,'Recorded native campaign / 4 October 2026',9,MUTED)
box(s,.8,4.46,4.4,.28,'PR / patch → test record → reproducible next step',11,TEAL,True)

s=prs.slides[6]
content(s,'Useful to maintainers, integrators and testers.','Leave reviewable building blocks for the community to adopt and extend.',
 'Community Benefit & Continuation Story',7,60,
 'Maintainers can review bounded patches with native measurements instead of reverse engineering a broad demo. Integrators can reuse diagnostic resource adapters, receiver observation contracts, CAN lighting contracts and the managed testbench pattern. Test engineers can reproduce interruption and recovery with machine-readable verdicts and labelled historical evidence. Our next steps are concrete: validate and resolve remaining upstream gates on the selected PRs; coordinate the OpenBSW router issue and contribution; then have another contributor exercise the integration path. We propose reuse through a Blueprint-style integration pattern, but do not claim an accepted Blueprint or quantified engineering savings. The benefit is a practical review and reproduction path across projects.',
 [('Maintainer PR', 'https://github.com/eclipse-score/communication/pull/1335'),('Integration handover',BASE+'contributions/shared/docs/handover.md'),('Router continuation',BASE+'contributions/eclipse-openbsw/transport-router/README.md')], 'h6b2788cdbe65c0ea_1_164')
columns(s,[('MAINTAINERS','Review and adopt','Bounded fixes and tests with recorded baselines and limitations.'),('INTEGRATORS','Reuse the pattern','Diagnostic adapters, vehicle contracts and deployment examples.'),('TEST ENGINEERS','Repeat the evidence','Fault campaigns, recovery checks and inspectable verdict artifacts.')])
banner(s,'Next: resolve native gates → obtain project review → repeat with another contributor.')

s=prs.slides[7]
content(s,'Existing project work, moved forward.','Each selected contribution has an issue, an artifact and a concrete next action.',
 'Contribution Focus & Initiative',5,60,
 'We deliberately connect our work to existing Eclipse project needs. Communication 1167 moved from a test gap to submitted PR 1335; the next step is native CI and maintainer review. Diagnostics 16 moved to submitted PR 40 with an OpenSOVD provider adapter; remaining author agreement and functional validation are explicit. Lifecycle 704 has a measured configuration-generation patch with 113 native passing cases and a prepared upstream submission. OpenBSW logical-address routing emerged from our integration need and has a verified local module plus an issue draft for approach agreement. Jefferson owns architecture, factory work and ECU contributions; Bruno owns S-CORE and X-Verse integration; Puru owns diagnostic adapters and the dashboard path; Yasser supports OpenSOVD, review and documentation; Siva supports openDuT and test evidence. Prepared core work dates to 4 October; the hardware and gateway have event-time results on 6–7 October. The final event delta and the actual day-one declaration must be recorded, rather than inferred.',
 [('Existing COM issue','https://github.com/eclipse-score/communication/issues/1167'),('Existing diagnostics issue','https://github.com/eclipse-score/inc_diagnostics/issues/16'),('Lifecycle #704',BASE+'contributions/eclipse-score/lifecycle/704/README.md')], 'h6b2788cdbe65c0ea_1_171')
columns(s,[('COM #1167','PR #1335','Integration coverage submitted. Next: CI and maintainer review.'),('DIAGNOSTICS #16','PR #40','Provider adapter submitted. Next: author/ECA and native validation.'),('LIFECYCLE #704','Verified patch','Configuration deduplication. Next: native upstream submission.')])
banner(s,'Prepared core: 4 Oct. Hardware and gateway evidence: 6–7 Oct. Attribution stays explicit.')

# Add a bonus slide from the inspected content exemplar; keep original closing.
s=clone(prs.slides[1])
# Clone a populated exemplar, then explicitly delete only destination-irrelevant added objects.
for sh in list(s.shapes):
    if sh.shape_id>79:
        sh._element.getparent().remove(sh._element)
content(s,'Technology evidence, ready for inspection.','Three bonus technologies have meaningful runtime or testbench evidence.',
 'Extra Technology Points',None,20,
 'openDuT is used in the managed network and disturbance campaign. ThreadX is actual runtime code in both Linux simulation and the AZ3166 embedded target. AutoSD is a real QEMU/KVM deployment target with build, diagnostics, interruption, recovery and reboot evidence. Under the official scorecard each meaningful technology can add 0.10, subject to evaluator acceptance and the final score cap. This supports asking evaluators to inspect three technologies, not claiming a bonus already awarded. Java/Jakarta has not been established as a qualifying implemented component by this review. No extra runtime was introduced merely as a slide label.',
 [('openDuT testbench',BASE+'contributions/eclipse-opendut/OpenDut/README.md'),('ThreadX hardware',BASE+'demo/X-Verse/external_hackathon_ecus/ThreadX/artifacts/az3166-uart-can/results.json'),('AutoSD VM',BASE+'contributions/eclipse-autosd/AutoSD/artifacts/vehicle-computer/results.json')], 'h6b2788cdbe65c0ea_1_82')
columns(s,[('openDuT / +0.10','Managed bench','Real receiver traffic, communication interruption and recovery.'),('THREADX / +0.10','Actual RTOS','Linux simulation and physical AZ3166 lighting-controller evidence.'),('AUTOSD / +0.10','Real guest','Deployment, diagnosis, managed disturbance and reboot evidence.')])
banner(s,'Potential +0.30, subject to evaluation. Java/Jakarta is not claimed; final score is capped at 5.00.')

closing=prs.slides[8]
text(next(x for x in closing.shapes if x.shape_id==135),'Built to be continued.',28,'FFFFFF',False,'Arial')
text(next(x for x in closing.shapes if x.shape_id==144),'Review the patches.\nRun the evidence.',18,'FFFFFF',True,'Inter')
closing_link=link(closing,.83,4.01,5.2,'github.com/Eclipse-SDV-Hackathon-Chapter-Four/Thinking_CAPs',REPO)
closing_link.text_frame.paragraphs[0].runs[0].hyperlink.address=None
closing_link.click_action.hyperlink.address=REPO
style(closing_link.text_frame.paragraphs[0].runs[0],10,'FFFFFF')
register(closing,'Built to be continued.','Closing',None,10,
 'Our handover is code, evidence and a next step. Review the selected contributions, repeat the campaign and help make this integration easier for the next Eclipse SDV contributor. Thank you.',[('Repository',REPO)],'h6b2788cdbe65c0ea_1_65')

# Put the bonus immediately before the native closing. Native slide IDs remain unique.
ids=prs.slides._sldIdLst
bonus_id=ids[-1]
ids.remove(bonus_id);ids.insert(8,bonus_id)
# plan registration had closing after bonus, so matches the delivered order.

# Appendix: all selected sources, without adding them to the timed pitch.
s=clone(prs.slides[1])
for sh in list(s.shapes):
    if sh.shape_id>79: sh._element.getparent().remove(sh._element)
content(s,'Evidence index for the technical interview.','Each link identifies a bounded contribution or a recorded integration result.',
 'Appendix / Evidence',None,0,
 'Use this slide for questions. The contribution and evidence links are pinned to inspected origin/main revision 38fbb58b. PR links are live and their status can change. Measurements are retained original records. The current deck was authored from inspected local work and origin/main; it is not an attestation that every latest working-tree update is published. Check the final submission revision before presenting.',[], 'h6b2788cdbe65c0ea_1_82')
sources=[('COM #1167 / native evidence',BASE+'contributions/eclipse-score/communication/1167/README.md'),('Lifecycle #704 / 113 cases',BASE+'contributions/eclipse-score/lifecycle/704/README.md'),('OpenBSW router / native gates',BASE+'contributions/eclipse-openbsw/transport-router/validation.md'),('Receiver / physical verdicts',BASE+'contributions/shared/evidence/f009-reproduction-physical/results.json'),('Dashboard / API and browser checks',BASE+'contributions/shared/evidence/f010-live/verification.json'),('ThreadX / hardware and CARLA',BASE+'demo/X-Verse/external_hackathon_ecus/ThreadX/artifacts/az3166-xverse-carla/results.json'),('AutoSD / guest and source identities',BASE+'contributions/eclipse-autosd/AutoSD/artifacts/README.md'),('Fault storage / regression and patch',BASE+'contributions/eclipse-opensovd/OpenSOVD/contributions/fault-storage-write-through/README.md')]
for i,(label,url) in enumerate(sources):
    x=.8+(i%2)*4.35;y=2.04+(i//2)*.60
    box(s,x,y,4.15,.28,label,14,PURPLE,True)
    link(s,x,y+.31,4.15,'Open source / evidence record',url)
plan[-1]['sources']=[dict(label=l,url=u) for l,u in sources]
s.notes_slide.notes_text_frame.text += '\n'+'\n'.join(l+': '+u for l,u in sources)

s=clone(prs.slides[1])
for sh in list(s.shapes):
    if sh.shape_id>79: sh._element.getparent().remove(sh._element)
content(s,'Team, scope and template attribution.','Five contributors; clear ownership; preserved preparation and event-work boundaries.',
 'Appendix / Team & Sources',None,0,
 'The official native template is [EF_SDV_Hackathon_2026] - Pitching Slides Template, linked from page five of the pitching-session PDF. We retained its cover, closing, branding, white content family and copyright footers, and adapted its narrative content for Freestyle. The source attribution and CC-BY-SA 4.0 notice are preserved. This adaptation is also made available under CC-BY-SA 4.0. Prepared-code disclosure is a separate organizer record; this slide does not attest it. The pitch target is 9 minutes 30 seconds; appendix slides are for questions.',
 [('Official slide template',TEMPLATE),('Pitching-session guide',GUIDE),('Freestyle scorecard',RUBRIC)], 'h6b2788cdbe65c0ea_1_82')
box(s,.8,2.10,8.5,2.40,'Jefferson Nascimento — architecture, factory, ECU contributions\nBruno Campos — S-CORE application and X-Verse integration\nPuru — OpenSOVD / CDA adapters and dashboard\nYasser — OpenSOVD, review and documentation\nSiva — openDuT and integration-test evidence',16)
banner(s,'Template: Eclipse Foundation AISBL and contributors / adapted under CC-BY-SA 4.0.')

prs.core_properties.title='Thinking CAPs — Eclipse SDV Freestyle Pitch 2026'
prs.core_properties.subject='One slide per Freestyle scoring criterion; evidence-linked contribution handover'
prs.core_properties.author='Thinking CAPs'
prs.core_properties.keywords='Eclipse SDV, Freestyle, HackFest, S-CORE, OpenSOVD, openDuT, ThreadX, AutoSD, OpenBSW'
prs.save(ROOT / 'Thinking_CAPs_Freestyle_Pitch_2026.pptx')
timed=sum(x['seconds'] for x in plan)
assert timed==570, timed
assert sum(x['weight_percent'] or 0 for x in plan)==100
assert len(prs.slides)==len(plan)==12
for i,item in enumerate(plan): item['slide_number']=i+1
(ROOT/'slide-plan.json').write_text(json.dumps(dict(pitch_seconds=timed,source_revision='38fbb58b4420365a8c02de1552ed983d51e9a499',template_url=TEMPLATE,slides=plan),indent=2)+'\n')
out=['THINKING CAPs — FREESTYLE PITCH','Target: 9:30. Slides 11–12 are Q&A appendices.','Native test figures are retained measurements; status snapshot: 7 October 2026.','']
elapsed=0
for i,(title,criterion,seconds,speech,sources) in enumerate(notes):
    window=f'{elapsed//60}:{elapsed%60:02d}–{(elapsed+seconds)//60}:{(elapsed+seconds)%60:02d}' if seconds else 'Q&A appendix'
    out.extend([f'SLIDE {i+1}: {title}',f'{criterion} | {window}',speech,'']+[label+': '+url for label,url in sources]+[''])
    elapsed+=seconds
(ROOT/'speaker-notes.txt').write_text('\n'.join(out)+'\n')
(ROOT/'template-attribution.txt').write_text('Source template: [EF_SDV_Hackathon_2026] - Pitching Slides Template\n'+TEMPLATE+'\nDiscovered in page 5 of: '+GUIDE+'\nCopyright © Eclipse Foundation AISBL and contributors.\nTemplate and this presentation adaptation: CC-BY-SA 4.0 International.\nhttps://creativecommons.org/licenses/by-sa/4.0/\nChanges: replaced instructional narrative, added Freestyle criterion labels, evidence links, diagrams, screenshots and speaker notes; retained source cover/closing artwork, branding and footers.\nEvidence URLs are pinned to inspected origin/main 38fbb58b. Native results are historical measurements. Prepared core and event work remain separately attributed.\n')
print(json.dumps(dict(slides=len(prs.slides),pitch_seconds=timed,pptx=str(ROOT/'Thinking_CAPs_Freestyle_Pitch_2026.pptx'))))
